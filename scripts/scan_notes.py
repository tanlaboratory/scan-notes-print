#!/usr/bin/env python3
"""Deterministic scanned-note enhancement; no OCR or generative reconstruction."""
import argparse
import hashlib
import json
import math
from pathlib import Path
import re
import shutil
import subprocess
import sys
from zipfile import ZipFile, ZIP_DEFLATED

import numpy as np
from PIL import Image, ImageDraw, ImageFilter, ImageFont, ImageOps

PRESETS = {'gentle': (.34, .91, 85), 'balanced': (.44, .85, 115), 'crisp': (.48, .80, 150)}
NS = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main',
      'a': 'http://schemas.openxmlformats.org/drawingml/2006/main',
      'wp': 'http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing',
      'r': 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'}


def run(cmd):
    result = subprocess.run([str(x) for x in cmd], capture_output=True, text=True)
    if result.returncode:
        raise RuntimeError((result.stderr or result.stdout)[-3000:])
    return result


def source_signature(paths):
    return [hashlib.sha256(Path(p).read_bytes()).hexdigest() for p in paths]


def extract(inputs, work, pdftoppm, dpi):
    raw = work / 'source'; raw.mkdir(parents=True, exist_ok=True)
    pages = []; docx = None
    if any(p.suffix.lower() == '.docx' for p in inputs) and len(inputs) != 1:
        raise ValueError('DOCX must be the only input.')
    for source_index, path in enumerate(inputs, 1):
        ext = path.suffix.lower()
        if ext == '.docx':
            from lxml import etree
            docx = path
            with ZipFile(path) as z:
                rels = etree.fromstring(z.read('word/_rels/document.xml.rels'))
                relmap = {r.get('Id'): r.get('Target') for r in rels}
                root = etree.fromstring(z.read('word/document.xml'))
                ordered = []
                for blip in root.findall('.//a:blip', NS):
                    target = relmap.get(blip.get('{%s}embed' % NS['r']))
                    if target and target not in ordered:
                        ordered.append(target)
                # Include header/footer media too, without interpreting document text.
                for name in sorted(z.namelist()):
                    if name.startswith('word/media/') and not name.endswith('/'):
                        target = name[5:]
                        if target not in ordered: ordered.append(target)
                for target in ordered:
                    name = 'word/' + target if not target.startswith('/') else target.lstrip('/')
                    if name not in z.namelist():
                        raise ValueError('Unsupported external/relative image relationship: ' + target)
                    ident = str(len(pages) + 1)
                    output = raw / (ident + Path(name).suffix)
                    output.write_bytes(z.read(name))
                    try:
                        with Image.open(output) as im:
                            im.load(); size = im.size
                    except Exception as exc:
                        raise ValueError('Unsupported embedded image: ' + name) from exc
                    pages.append({'id': ident, 'path': str(output), 'media': name,
                                  'size': list(size), 'source': str(path)})
        elif ext == '.pdf':
            tool = pdftoppm or shutil.which('pdftoppm')
            if not tool: raise ValueError('pdftoppm is required; pass --pdftoppm from bundled dependencies.')
            prefix = raw / f'pdf-{source_index}'
            for old in raw.glob(prefix.name + '-*.png'): old.unlink()
            run([tool, '-r', str(dpi), '-png', path, prefix])
            matches = sorted(raw.glob(prefix.name + '-*.png'),
                             key=lambda p: int(p.stem.rsplit('-', 1)[1]))
            if not matches: raise ValueError('PDF rendering produced no pages.')
            for p in matches:
                with Image.open(p) as im: size = im.size
                pages.append({'id': str(len(pages)+1), 'path': str(p), 'size': list(size), 'source': str(path)})
        else:
            with Image.open(path) as im:
                for frame in range(getattr(im, 'n_frames', 1)):
                    im.seek(frame)
                    img = ImageOps.exif_transpose(im.copy()).convert('RGB')
                    ident = str(len(pages)+1); output = raw / f'{ident}.png'; img.save(output)
                    pages.append({'id': ident, 'path': str(output), 'size': list(img.size), 'source': str(path)})
    if not pages: raise ValueError('No processable images found.')
    return pages, docx


def contact_sheet(pages, output, key='path'):
    if len(pages) > 16:
        files = []
        for start in range(0, len(pages), 16):
            part = output.with_name(output.stem + '-' + str(start//16 + 1) + output.suffix)
            files.extend(contact_sheet(pages[start:start+16], part, key))
        return files
    cols = 4; cellw, cellh = 280, 350
    sheet = Image.new('RGB', (cols*cellw, math.ceil(len(pages)/cols)*cellh), 'white')
    draw = ImageDraw.Draw(sheet)
    for n, p in enumerate(pages):
        with Image.open(p[key]) as im:
            im = im.convert('RGB'); im.thumbnail((cellw-16, cellh-32), Image.Resampling.LANCZOS)
            x, y = (n%cols)*cellw, (n//cols)*cellh
            sheet.paste(im, (x+(cellw-im.width)//2, y+26))
            draw.text((x+8,y+5), f"#{p['id']}   {p['size'][0]} x {p['size'][1]}", fill='black')
    sheet.save(output)
    return [str(output)]


def enhance(path, output, crop=None, mode='auto', preset='balanced', monochrome=False):
    with Image.open(path) as opened:
        im = ImageOps.exif_transpose(opened).convert('RGB')
    original_size = im.size
    if crop is not None:
        if len(crop) != 4 or not all(isinstance(x, (int,float)) and math.isfinite(x) for x in crop):
            raise ValueError('Crop must contain four finite numbers.')
        l,t,r,b = crop
        if not (0 <= l < r <= 1 and 0 <= t < b <= 1):
            raise ValueError('Crop must satisfy 0 <= left < right <= 1 and 0 <= top < bottom <= 1.')
        box = tuple(round(v*d) for v,d in zip(crop, (im.width,im.height,im.width,im.height)))
        if box[2]-box[0] < 32 or box[3]-box[1] < 32: raise ValueError('Crop is too small.')
        im = im.crop(box)
    a = np.asarray(im, dtype=np.float32)
    if mode == 'auto':
        # A white screenshot/clean scan usually has a large exact-white population.
        mode = 'preserve' if np.mean(np.min(a,axis=2) > 247) > .55 else 'pen'
    low, high, sharpness = PRESETS[preset]
    if mode == 'pen':
        # Per-channel local illumination correction also removes paper color casts.
        small = im.resize((max(16,im.width//8), max(16,im.height//8)), Image.Resampling.BILINEAR)
        channels = [c.filter(ImageFilter.MaxFilter(15)).filter(ImageFilter.GaussianBlur(4))
                    for c in small.split()]
        bg = Image.merge('RGB',channels).resize(im.size,Image.Resampling.BICUBIC)
        corrected = np.clip(a / np.maximum(np.asarray(bg,dtype=np.float32),1),0,1)
        darkest = corrected.min(axis=2)
        density = 1-np.clip((darkest-low)/(high-low),0,1)
        chroma = corrected.max(axis=2)-darkest
        strength = np.clip((chroma-.045)/.10,0,1)
        hue = (corrected-darkest[:,:,None])/np.maximum(chroma[:,:,None],.001)
        ink = 185*hue*strength[:,:,None]
        if monochrome: ink[:] = 0
        clean = 255-density[:,:,None]*(255-ink)
    else:
        # Preserve colored diagrams/highlights rather than interpreting them as pen strokes.
        clean = np.clip((a-25)/(245-25),0,1)*255
        if monochrome:
            g = clean[:,:,0]*.299+clean[:,:,1]*.587+clean[:,:,2]*.114
            clean = np.repeat(g[:,:,None],3,axis=2)
    result = Image.fromarray(clean.astype('uint8'))
    # Interpolation improves edge smoothness at print scale; it adds no new detail.
    scale = min(2.0, math.sqrt(16_000_000/(im.width*im.height)))
    if scale > 1.01:
        result = result.resize((round(im.width*scale),round(im.height*scale)),Image.Resampling.LANCZOS)
    result = result.filter(ImageFilter.UnsharpMask(radius=1.1,percent=sharpness,threshold=4))
    result.save(output, dpi=(300,300))
    pixels = np.asarray(result)
    return {'mode':mode,'crop':crop,'original_size':list(original_size),'size':list(result.size),
            'white_fraction':round(float(np.mean(pixels.min(axis=2)>=250)),4)}


def patch_docx(source, pages, output):
    from lxml import etree
    with ZipFile(source) as z: files = {n:z.read(n) for n in z.namelist()}
    replacements = {p['media']: f"word/media/print-note-{p['id']}.png" for p in pages}
    by_media = {p['media']:p for p in pages}
    for name in list(files):
        if not name.endswith('.rels'): continue
        root = etree.fromstring(files[name]); ridmap = {}; changed = False
        for el in root:
            target = el.get('Target','')
            media = 'word/' + target if not target.startswith('/') else target.lstrip('/')
            if media in replacements:
                el.set('Target',replacements[media][5:]); ridmap[el.get('Id')] = by_media[media];changed=True
        if not changed: continue
        files[name] = etree.tostring(root,xml_declaration=True,encoding='UTF-8',standalone=True)
        # For standard word/_rels/document.xml.rels and header/footer equivalents.
        owner = str(Path(name).parent.parent / Path(name).name.removesuffix('.rels'))
        if owner not in files: continue
        xml = etree.fromstring(files[owner])
        for inline in xml.findall('.//wp:inline',NS):
            blip=inline.find('.//a:blip',NS)
            if blip is None: continue
            p=ridmap.get(blip.get('{%s}embed'%NS['r']))
            if not p or p.get('crop') is None: continue
            ext=inline.find('wp:extent',NS)
            if ext is None: continue
            oldw,oldh=int(ext.get('cx')),int(ext.get('cy'))
            w,h=p['size']; width=min(oldw,round(oldh*w/h));height=round(width*h/w)
            ext.set('cx',str(width));ext.set('cy',str(height))
            for e in inline.findall('.//a:xfrm/a:ext',NS):e.set('cx',str(width));e.set('cy',str(height))
        files[owner] = etree.tostring(xml,xml_declaration=True,encoding='UTF-8',standalone=True)
    ct=etree.fromstring(files['[Content_Types].xml'])
    if not any(e.get('Extension')=='png' for e in ct):
        etree.SubElement(ct,'{http://schemas.openxmlformats.org/package/2006/content-types}Default',Extension='png',ContentType='image/png')
    files['[Content_Types].xml']=etree.tostring(ct,xml_declaration=True,encoding='UTF-8',standalone=True)
    with ZipFile(output,'w',ZIP_DEFLATED) as z:
        for n,v in files.items():
            if n not in replacements:z.writestr(n,v)
        for p in pages:z.write(p['enhanced'],replacements[p['media']])


def safe_text_render_copy(source, output, font_path):
    from docx import Document
    from docx.shared import Inches
    from docx.oxml.ns import qn
    d=Document(source)
    if d.tables or len(d.sections)!=1: raise ValueError('Print-safe text only supports a simple single-section DOCX without tables.')
    if d._element.findall('.//'+qn('wp:anchor')): raise ValueError('Print-safe text does not support floating images.')
    for section in d.sections:
        if section.header._element.findall('.//'+qn('w:t')) or section.footer._element.findall('.//'+qn('w:t')):
            raise ValueError('Print-safe text does not support text in headers/footers.')
    if d._element.findall('.//'+qn('w:txbxContent')) or d._element.findall('.//'+qn('m:oMath')):
        raise ValueError('Print-safe text does not support text boxes or native equations.')
    font=ImageFont.truetype(str(font_path),36)
    directory=output.parent/'text';directory.mkdir(exist_ok=True)
    section=d.sections[0];inches=(section.page_width-section.left_margin-section.right_margin)/914400
    pixels=round(inches*200)
    for n,p in enumerate(d.paragraphs):
        if not p.text:continue
        if p._element.findall('.//'+qn('w:drawing')):
            raise ValueError('Print-safe text cannot replace mixed text/image paragraphs.')
        text=p.text.replace('➕','+').replace('\ufe0f','').replace('\u20e3','')
        lines=[];current=''
        for ch in text:
            if ch=='\n' or font.getlength(current+ch)>pixels-12:
                lines.append(current);current=''
                if ch=='\n':continue
            current+=ch
        lines.append(current)
        im=Image.new('RGB',(pixels,len(lines)*52+10),'white');draw=ImageDraw.Draw(im)
        for k,line in enumerate(lines):draw.text((3,k*52+3),line,font=font,fill='black')
        path=directory/f'{n}.png';im.save(path)
        p.clear()  # pPr, including list numbering, is retained.
        p.add_run().add_picture(str(path),width=Inches(inches))
    d.save(output)


def image_pdf(pages, output, paper, orientation, margin):
    from reportlab.pdfgen import canvas
    from reportlab.lib.pagesizes import A4,letter
    from reportlab.lib.utils import ImageReader
    base=A4 if paper=='A4' else letter
    c=canvas.Canvas(str(output),pagesize=base,pageCompression=1)
    c.setTitle('Clear color print notes')
    for p in pages:
        w,h=p['size']
        landscape=orientation=='landscape' or (orientation=='auto' and w>h)
        pw,ph=base[::-1] if landscape else base
        if margin*2>=min(pw,ph):raise ValueError('Margins leave no printable page area.')
        c.setPageSize((pw,ph))
        scale=min((pw-margin*2)/w,(ph-margin*2)/h)
        dw,dh=w*scale,h*scale
        c.drawImage(ImageReader(p['enhanced']),(pw-dw)/2,(ph-dh)/2,width=dw,height=dh)
        c.showPage()
    c.save()


def main():
    ap=argparse.ArgumentParser(description=__doc__)
    ap.add_argument('command',choices=['inspect','build'])
    ap.add_argument('--input',action='append',required=True)
    ap.add_argument('--work-dir',required=True)
    ap.add_argument('--output-dir')
    ap.add_argument('--name')
    ap.add_argument('--config')
    ap.add_argument('--preset',choices=list(PRESETS),default='balanced')
    ap.add_argument('--monochrome',action='store_true')
    ap.add_argument('--paper',choices=['A4','Letter'],default='A4')
    ap.add_argument('--orientation',choices=['portrait','landscape','auto'],default='portrait')
    ap.add_argument('--margin-mm',type=float,default=10)
    ap.add_argument('--dpi',type=int,default=240)
    ap.add_argument('--pdftoppm')
    ap.add_argument('--renderer')
    ap.add_argument('--print-safe-text',action='store_true')
    ap.add_argument('--font')
    ap.add_argument('--overwrite',action='store_true')
    args=ap.parse_args()
    if args.dpi<72 or args.dpi>600:ap.error('--dpi must be 72..600.')
    if not math.isfinite(args.margin_mm) or args.margin_mm<0:ap.error('Invalid margin.')
    if args.command=='build' and not args.output_dir:ap.error('--output-dir required for build.')
    paths=[Path(p).expanduser().resolve() for p in args.input]
    for p in paths:
        if not p.is_file():raise ValueError('Input file not found: '+str(p))
    work=Path(args.work_dir).resolve();work.mkdir(parents=True,exist_ok=True)
    signature=source_signature(paths)
    cache=work/'manifest.json'
    saved=json.loads(cache.read_text()) if cache.exists() else {}
    if saved.get('signature')==signature and saved.get('dpi')==args.dpi and all(Path(p['path']).exists() for p in saved.get('pages',[])):
        pages=saved['pages'];docx=Path(saved['docx']) if saved.get('docx') else None
    else:
        pages,docx=extract(paths,work,args.pdftoppm,args.dpi)
        cache.write_text(json.dumps({'signature':signature,'dpi':args.dpi,'pages':pages,'docx':str(docx) if docx else None},ensure_ascii=False,indent=2))
    source_sheet=work/'source-contact.jpg';source_sheets=contact_sheet(pages,source_sheet)
    if args.command=='inspect':
        print(json.dumps({'images':len(pages),'contact_sheets':source_sheets,'manifest':str(cache),
                          'pages':[{'id':p['id'],'size':p['size'],'path':p['path'],'media':p.get('media')} for p in pages[:20]],'more_pages':max(0,len(pages)-20)},ensure_ascii=False))
        return
    config=json.loads(Path(args.config).read_text()) if args.config else {}
    ids={p['id'] for p in pages}
    if not set(config.get('crops',{})).issubset(ids) or not set(config.get('modes',{})).issubset(ids):
        raise ValueError('Config contains unknown image IDs; inspect this input first.')
    out=Path(args.output_dir).resolve();out.mkdir(parents=True,exist_ok=True)
    name=args.name or (paths[0].stem+'_彩色清晰打印版' if not args.monochrome else paths[0].stem+'_黑白清晰打印版')
    if Path(name).name!=name or name in ('.','..'):raise ValueError('--name must be a filename stem, not a path.')
    pdf=out/(name+'.pdf');docout=out/(name+'.docx')
    for p in ([pdf,docout] if docx else [pdf]):
        if p.exists() and not args.overwrite:raise ValueError('Output exists; use --overwrite for this task output: '+str(p))
        if p in paths:raise ValueError('Cannot overwrite an input.')
    if docx and not args.renderer:raise ValueError('DOCX requires --renderer pointing to documents/render_docx.py.')
    clean=work/'enhanced';clean.mkdir(exist_ok=True)
    for p in pages:
        p['enhanced']=str(clean/(p['id']+'.png'))
        mode=config.get('modes',{}).get(p['id'],'auto')
        if mode not in ('auto','pen','preserve'):raise ValueError('Unknown mode: '+mode)
        p.update(enhance(p['path'],p['enhanced'],config.get('crops',{}).get(p['id']),mode,args.preset,args.monochrome))
    sheet=work/'enhanced-contact.jpg';sheets=contact_sheet(pages,sheet,'enhanced')
    if docx:
        patch_docx(docx,pages,docout)
        render_input=docout
        if args.print_safe_text:
            font=args.font
            if not font:raise ValueError('--font required for --print-safe-text.')
            render_input=work/'print-safe.docx';safe_text_render_copy(docout,render_input,font)
        render=work/'render';run([sys.executable,args.renderer,render_input,'--output_dir',render,'--emit_pdf'])
        rendered=render/(render_input.stem+'.pdf')
        if not rendered.exists():raise ValueError('Renderer did not create a PDF.')
        shutil.copy2(rendered,pdf)
    else:
        image_pdf(pages,pdf,args.paper,args.orientation,args.margin_mm*72/25.4)
    from pypdf import PdfReader
    count=len(PdfReader(pdf).pages)
    if source_signature(paths)!=signature:raise RuntimeError('Input changed during processing.')
    report={'images':len(pages),'pdf_pages':count,'pdf':str(pdf),'docx':str(docout) if docx else None,
            'enhanced_contacts':sheets,'render_dir':str(work/'render') if docx else None,
            'preset':args.preset,'monochrome':args.monochrome,
            'pages':[{k:p.get(k) for k in ('id','mode','crop','size','white_fraction')} for p in pages]}
    (work/'report.json').write_text(json.dumps(report,ensure_ascii=False,indent=2))
    print(json.dumps({k:v for k,v in report.items() if k!='pages'} | {'report':str(work/'report.json')},ensure_ascii=False))


if __name__=='__main__':
    try:main()
    except (ValueError,RuntimeError,OSError) as exc:
        print(json.dumps({'error':str(exc)},ensure_ascii=False),file=sys.stderr);sys.exit(1)
